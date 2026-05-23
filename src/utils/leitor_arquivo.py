from __future__ import annotations

import io
from pathlib import Path
from typing import Union

import fitz
from docx import Document


class LeitorArquivo:
    """Lê e extrai texto de arquivos PDF, DOCX e TXT."""

    FORMATOS_SUPORTADOS = {".pdf", ".docx", ".txt"}

    def ler(
        self,
        origem: Union[str, Path, bytes, io.BytesIO],
        nome_arquivo: str = "",
    ) -> str:

        extensao = self._obter_extensao(origem, nome_arquivo)

        if extensao == ".pdf":
            return self._ler_pdf(origem)

        elif extensao == ".docx":
            return self._ler_docx(origem)

        elif extensao == ".txt":
            return self._ler_txt(origem)

        raise ValueError(
            f"Formato '{extensao}' não suportado. "
            f"Use: {', '.join(self.FORMATOS_SUPORTADOS)}"
        )

    def _obter_extensao(
        self,
        origem: Union[str, Path, bytes, io.BytesIO],
        nome_arquivo: str,
    ) -> str:

        if isinstance(origem, (str, Path)):
            return Path(origem).suffix.lower()

        return Path(nome_arquivo).suffix.lower() if nome_arquivo else ""

    def _ler_pdf(self, origem) -> str:
        """Extrai texto de PDF usando PyMuPDF."""

        doc = None

        try:
            if isinstance(origem, (str, Path)):
                doc = fitz.open(str(origem))

            elif isinstance(origem, bytes):
                doc = fitz.open(stream=origem, filetype="pdf")

            else:
                origem.seek(0)
                doc = fitz.open(
                    stream=origem.read(),
                    filetype="pdf"
                )

            paginas = [
                pagina.get_text()
                for pagina in doc
            ]

            return "\n".join(paginas)

        except Exception as e:
            raise IOError(f"Erro ao ler PDF: {e}") from e

        finally:
            if doc:
                doc.close()

    def _ler_docx(self, origem) -> str:
        """Extrai texto de DOCX."""

        try:

            if isinstance(origem, (str, Path)):
                doc = Document(str(origem))

            elif isinstance(origem, bytes):
                doc = Document(io.BytesIO(origem))

            else:
                origem.seek(0)
                doc = Document(origem)

            paragrafos = [
                p.text.strip()
                for p in doc.paragraphs
                if p.text.strip()
            ]

            for tabela in doc.tables:
                for linha in tabela.rows:
                    for celula in linha.cells:

                        texto = celula.text.strip()

                        if texto:
                            paragrafos.append(texto)

            return "\n".join(paragrafos)

        except Exception as e:
            raise IOError(f"Erro ao ler DOCX: {e}") from e

    def _ler_txt(self, origem) -> str:
        """Lê TXT."""

        try:

            if isinstance(origem, (str, Path)):
                return Path(origem).read_text(
                    encoding="utf-8",
                    errors="ignore",
                )

            elif isinstance(origem, bytes):
                return origem.decode(
                    "utf-8",
                    errors="ignore",
                )

            else:
                origem.seek(0)

                return origem.read().decode(
                    "utf-8",
                    errors="ignore",
                )

        except Exception as e:
            raise IOError(f"Erro ao ler TXT: {e}") from e

    def formato_suportado(self, nome_arquivo: str) -> bool:
        return (
            Path(nome_arquivo)
            .suffix.lower()
            in self.FORMATOS_SUPORTADOS
        )


leitor = LeitorArquivo()


def ler_arquivo(
    origem: Union[str, Path, bytes, io.BytesIO],
    nome_arquivo: str = "",
) -> str:
    return leitor.ler(origem, nome_arquivo)